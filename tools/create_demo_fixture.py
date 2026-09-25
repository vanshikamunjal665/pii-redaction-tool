"""Create a clearly labelled synthetic DOCX fixture for reproducible testing.

The assignment attachment was not present in the execution environment.  This
fixture is not presented as an RHP; it exists so the complete pipeline and its
reports can be exercised without using a real person's document.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = Path(__file__).resolve().parents[1] / "input" / "synthetic_rhp_fixture.docx"


def add_hyperlink(paragraph, text: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction_node, separate, text, end])


def build() -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    document.core_properties.title = "Synthetic RHP Fixture - Demonstration Only"
    document.core_properties.subject = "Synthetic PII detector test fixture"
    document.core_properties.author = "Synthetic Fixture Generator"

    title = document.add_heading("Synthetic Red Herring Prospectus Fixture", level=0)
    title.alignment = 1
    subtitle = document.add_paragraph("DEMONSTRATION ONLY — all values below are fabricated test data.")
    subtitle.runs[0].bold = True

    document.add_heading("Issuer and contact details", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Issuer", "Acme Technologies Pvt. Ltd."),
        ("Contact person", "Rashi Patil"),
        ("Registered office", "221B, Green Park, New Delhi, Delhi 110001"),
        ("Email", "rashi.patil@gmail.com"),
        ("Telephone", "+91 98765 43210"),
        ("Date of Birth", "14/08/1998"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_paragraph(
        "The issuer, Acme Technologies Pvt. Ltd., confirms that Rashi Patil is the contact person."
    )
    document.add_paragraph(
        "Registered office: 221B, Green Park, New Delhi, Delhi 110001; telephone +91-98765-43210."
    )
    document.add_paragraph("Email: ")
    add_hyperlink(document.paragraphs[-1], "rashi.patil@gmail.com", "mailto:rashi.patil@gmail.com")

    document.add_heading("Synthetic validation fields", level=1)
    synthetic_table = document.add_table(rows=1, cols=2)
    synthetic_table.style = "Light Shading Accent 1"
    synthetic_table.rows[0].cells[0].text = "Field"
    synthetic_table.rows[0].cells[1].text = "Value"
    for label, value in [
        ("Social Security Number", "123-45-6789"),
        ("Credit card", "4111 1111 1111 1111"),
        ("IP address", "192.168.1.10"),
        ("Backup organisation", "Northstar Foods Limited"),
    ]:
        cells = synthetic_table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_heading("Signatories", level=1)
    document.add_paragraph("Director: John Doe", style="List Bullet")
    document.add_paragraph("Signatory: Rashi Patil", style="List Bullet")
    document.add_paragraph("Authorised signatory: John Doe", style="List Bullet")

    document.add_heading("Non-PII controls", level=1)
    for item in [
        "Order 123456789",
        "Invoice 1234567890123456",
        "Revenue 2025",
        "Page 123",
        "Ticket #12345",
        "Filing date: 12/01/2025",
        "Transaction date: 01/02/2026",
    ]:
        document.add_paragraph(item, style="List Bullet")

    header = section.header.paragraphs[0]
    header.text = "Acme Technologies Pvt. Ltd. | +91 98765 43210"
    footer = section.footer.paragraphs[0]
    footer.text = "Confidential synthetic demonstration | Contact: rashi.patil@gmail.com | Page "
    add_field(footer, "PAGE")

    document.add_page_break()
    document.add_heading("Additional contact block", level=1)
    document.add_paragraph(
        "Notice may be sent to Rashi Patil at 221B, Green Park, New Delhi, Delhi 110001."
    )
    document.add_paragraph("Northstar Foods Limited is a separate synthetic organisation.")
    document.add_paragraph("Alternate phone: +1-202-555-0147")
    document.add_paragraph("IP endpoint: 192.168.1.10")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    print(build())
