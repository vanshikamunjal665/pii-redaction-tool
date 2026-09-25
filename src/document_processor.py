"""DOCX/PDF document processing with formatting-preserving substitutions."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Optional

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from .detectors import DOB_CONTEXT_RE, PIIDetector
from .models import Entity, PIIType, TextUnit
from .replacers import ReplacementRegistry

XML_NS = "http://www.w3.org/XML/1998/namespace"


@dataclass
class _DocxUnit:
    part: str
    unit_id: str
    element: Any
    text: str


@dataclass
class RedactionResult:
    predictions: list[Entity]
    source_units: list[TextUnit]
    redacted_units: list[TextUnit]
    output_path: Path

    def prediction_dicts(self) -> list[dict[str, Any]]:
        return [entity.as_dict() for entity in self.predictions]


def _part_name(part: Any) -> str:
    value = str(getattr(part, "partname", "unknown"))
    return value.lstrip("/")


def _text_segments(paragraph: Any) -> tuple[str, list[tuple[Any, int, int]]]:
    """Return visible text and editable ``w:t`` node ranges.

    Tabs and breaks are included in offsets so a detector never sees a false
    adjacency across a line break, while only text nodes are modified.
    """

    pieces: list[str] = []
    text_nodes: list[tuple[Any, int, int]] = []
    position = 0
    for node in paragraph.iter():
        if node.tag == qn("w:t"):
            value = node.text or ""
            pieces.append(value)
            text_nodes.append((node, position, position + len(value)))
            position += len(value)
        elif node.tag in {qn("w:tab"), qn("w:br"), qn("w:cr")}:
            pieces.append("\n" if node.tag != qn("w:tab") else "\t")
            position += 1
    return "".join(pieces), text_nodes


def _set_text_node(node: Any, value: str) -> None:
    node.text = value
    if value[:1].isspace() or value[-1:].isspace():
        node.set(f"{{{XML_NS}}}space", "preserve")
    else:
        node.attrib.pop(f"{{{XML_NS}}}space", None)


def _replace_span(paragraph: Any, start: int, end: int, replacement: str) -> bool:
    """Replace a character span while retaining the first run's formatting."""

    _, nodes = _text_segments(paragraph)
    overlapping = [(node, left, right) for node, left, right in nodes if left < end and start < right]
    if not overlapping:
        return False
    first_index = 0
    last_index = len(overlapping) - 1
    for index, (node, left, right) in enumerate(overlapping):
        local_start = max(0, start - left)
        local_end = min(right - left, end - left)
        current = node.text or ""
        prefix = current[:local_start]
        suffix = current[local_end:]
        if index == first_index:
            _set_text_node(node, prefix + replacement + (suffix if index == last_index else ""))
        else:
            _set_text_node(node, suffix if index == last_index else "")
    return True


def _iter_part_elements(document: Any) -> list[tuple[str, Any, Any]]:
    """Return unique document/header/footer/notes XML parts."""

    parts: list[tuple[str, Any, Any]] = []
    seen_parts: set[int] = set()
    seen_elements: set[int] = set()

    def add(part: Any) -> None:
        if part is None or id(part) in seen_parts:
            return
        seen_parts.add(id(part))
        element = getattr(part, "element", None)
        if element is None or id(element) in seen_elements:
            return
        seen_elements.add(id(element))
        parts.append((_part_name(part), part, element))

    add(document.part)
    for section in document.sections:
        for accessor in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            try:
                story = getattr(section, accessor)
                # Accessing a linked story can materialize a new empty
                # header/footer part in python-docx.  Only process stories that
                # are explicitly present in the source package.
                if getattr(story, "is_linked_to_previous", False):
                    continue
                add(story.part)
            except Exception:
                # python-docx can expose read-only/unsupported story parts
                # for some producer-specific documents.
                continue

    # Include notes/comments when the package exposes their XML parts.  This
    # is best effort and does not require optional python-docx APIs.
    try:
        for part in document.part.package.parts:
            name = _part_name(part)
            if name.endswith(("/footnotes.xml", "/endnotes.xml", "/comments.xml")):
                add(part)
    except Exception:
        pass
    return parts


def _collect_docx_units(document: Any) -> list[_DocxUnit]:
    units: list[_DocxUnit] = []
    for part_name, _part, element in _iter_part_elements(document):
        paragraph_index = 0
        for paragraph in element.iter(qn("w:p")):
            text, _ = _text_segments(paragraph)
            unit_id = f"{part_name}:p{paragraph_index:04d}"
            units.append(_DocxUnit(part_name, unit_id, paragraph, text))
            paragraph_index += 1
    return units


def _public_units(units: Iterable[_DocxUnit]) -> list[TextUnit]:
    return [TextUnit(unit.part, unit.unit_id, unit.text) for unit in units]


def extract_docx_units(path: str | Path) -> list[TextUnit]:
    """Extract paragraph/table/header/footer units from a DOCX."""

    document = Document(str(path))
    return _public_units(_collect_docx_units(document))


def extract_docx_text(path: str | Path) -> str:
    return "\n".join(unit.text for unit in extract_docx_units(path))


def _update_hyperlink_targets(
    document: Any, email_replacements: dict[str, str]
) -> None:
    if not email_replacements:
        return
    for _part_name_value, part, _element in _iter_part_elements(document):
        for relationship in list(getattr(part, "rels", {}).values()):
            if getattr(relationship, "reltype", "") != RT.HYPERLINK:
                continue
            target = getattr(relationship, "target_ref", "")
            updated = target
            for original, replacement in email_replacements.items():
                updated = re.sub(re.escape(original), replacement, updated, flags=re.IGNORECASE)
            if updated != target:
                # python-docx's relationship object stores the external target
                # in _target; assigning it preserves the relationship ID.
                relationship._target = updated


def redact_docx(
    input_path: str | Path,
    output_path: str | Path,
    detector: Optional[PIIDetector] = None,
    registry: Optional[ReplacementRegistry] = None,
) -> RedactionResult:
    """Redact a DOCX in place in a copied document and save it separately.

    Paragraph, table-cell, header/footer, hyperlink-display, and text-box
    ``w:p`` nodes are processed.  Existing styles and run properties are kept;
    when an entity crosses runs, the replacement is placed in the first run
    and the remaining original text is cleared.
    """

    source = Path(input_path)
    target = Path(output_path)
    if source.resolve() == target.resolve():
        raise ValueError("Input and output paths must be different")
    if source.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx input, got: {source}")

    detector = detector or PIIDetector()
    registry = registry or ReplacementRegistry()
    document = Document(str(source))
    units = _collect_docx_units(document)
    predictions: list[Entity] = []
    email_replacements: dict[str, str] = {}

    for index, unit in enumerate(units):
        # Table labels and values are separate paragraphs/units in DOCX.  A
        # narrow DOB label immediately before a value is supplied as context;
        # no arbitrary neighbouring text is used.
        context_prefix = ""
        if index > 0 and DOB_CONTEXT_RE.search(units[index - 1].text):
            context_prefix = units[index - 1].text
        entities = detector.detect(unit.text, unit.unit_id, context_prefix=context_prefix)
        # Apply right-to-left so offsets in the original unit remain valid.
        for entity in sorted(entities, key=lambda item: (item.start, item.end), reverse=True):
            replacement = registry.replacement_for(entity)
            if _replace_span(unit.element, entity.start, entity.end, replacement):
                predictions.append(entity)
                if entity.pii_type == PIIType.EMAIL:
                    email_replacements[entity.text] = replacement
    _update_hyperlink_targets(document, email_replacements)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))

    # Reopen explicitly so a malformed package fails during processing rather
    # than being reported as a successful run.
    reopened = Document(str(target))
    reopened_units = _collect_docx_units(reopened)
    return RedactionResult(
        predictions=predictions,
        source_units=_public_units(units),
        redacted_units=_public_units(reopened_units),
        output_path=target,
    )


def _pdf_page_units(pdf_path: str | Path) -> list[TextUnit]:
    try:
        import pdfplumber  # type: ignore
    except ImportError as exc:  # pragma: no cover - dependency error path
        raise RuntimeError("PDF input requires pdfplumber; install requirements.txt") from exc

    units: list[TextUnit] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                units.append(TextUnit("source", f"page-{page_number:03d}", text))
    return units


def redact_pdf(
    input_path: str | Path,
    output_path: str | Path,
    detector: Optional[PIIDetector] = None,
    registry: Optional[ReplacementRegistry] = None,
) -> RedactionResult:
    """Extract PDF text and create a readable DOCX containing redacted text.

    PDF layout is not reconstructed as vector-perfect DOCX.  Page boundaries,
    line breaks, and paragraph text are retained; this is safer than claiming
    exact formatting that pdfplumber cannot round-trip.
    """

    source = Path(input_path)
    target = Path(output_path)
    if source.resolve() == target.resolve():
        raise ValueError("Input and output paths must be different")
    if source.suffix.lower() != ".pdf":
        raise ValueError(f"Expected a .pdf input, got: {source}")

    detector = detector or PIIDetector()
    registry = registry or ReplacementRegistry()
    source_units = _pdf_page_units(source)
    document = Document()
    predictions: list[Entity] = []
    for unit in source_units:
        entities = detector.detect(unit.text, unit.unit_id)
        redacted, _ = registry.redact_text(unit.text, entities)
        predictions.extend(entities)
        document.add_heading(f"Source page {unit.unit_id.split('-')[-1]}", level=2)
        for line in redacted.splitlines() or [""]:
            document.add_paragraph(line)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    reopened = Document(str(target))
    return RedactionResult(
        predictions=predictions,
        source_units=source_units,
        redacted_units=_public_units(_collect_docx_units(reopened)),
        output_path=target,
    )


def redact_document(
    input_path: str | Path,
    output_path: str | Path,
    detector: Optional[PIIDetector] = None,
    registry: Optional[ReplacementRegistry] = None,
) -> RedactionResult:
    suffix = Path(input_path).suffix.lower()
    if suffix == ".docx":
        return redact_docx(input_path, output_path, detector, registry)
    if suffix == ".pdf":
        return redact_pdf(input_path, output_path, detector, registry)
    raise ValueError("Supported input formats are .docx and .pdf")
