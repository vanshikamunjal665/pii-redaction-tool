from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.detectors import PIIDetector
from src.document_processor import extract_docx_text, extract_docx_units, redact_docx
from src.replacers import ReplacementRegistry


def add_hyperlink(paragraph, text: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def test_docx_redaction_preserves_structure_and_updates_hyperlink(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    target = tmp_path / "redacted.docx"
    document = Document()
    document.add_heading("Issuer details", level=1)
    paragraph = document.add_paragraph("Contact: Rashi Patil; email: ")
    add_hyperlink(paragraph, "rashi.patil@gmail.com", "mailto:rashi.patil@gmail.com")
    split_name = document.add_paragraph("Authorised signatory: ")
    split_name.add_run("Rashi ")
    split_name.add_run("Patil")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Registered office"
    table.cell(0, 1).text = "221B, Green Park, New Delhi, Delhi 110001"
    table.cell(1, 0).text = "Date of Birth"
    table.cell(1, 1).text = "14/08/1998"
    document.sections[0].header.paragraphs[0].text = "Acme Technologies Pvt. Ltd."
    document.sections[0].footer.paragraphs[0].text = "Page 1"
    document.save(source)

    result = redact_docx(
        source,
        target,
        detector=PIIDetector(use_spacy=False),
        registry=ReplacementRegistry(seed=7),
    )
    assert target.is_file()
    assert result.predictions
    text = extract_docx_text(target)
    assert "Rashi Patil" not in text
    assert "rashi.patil@gmail.com" not in text
    assert "Acme Technologies Pvt. Ltd." not in text
    assert "14/08/1998" not in text
    assert "Issuer details" in text
    assert "Page 1" in text

    units = extract_docx_units(target)
    assert any("example.com" in unit.text for unit in units)
    # Reopening the package and inspecting relationships catches stale mailto
    # targets that a plain text-only implementation would miss.
    reopened = Document(str(target))
    targets = [
        rel.target_ref
        for rel in reopened.part.rels.values()
        if rel.reltype == RT.HYPERLINK
    ]
    assert targets
    assert all("rashi.patil@gmail.com" not in target for target in targets)
    assert all("example.com" in target for target in targets)


def test_redaction_refuses_to_overwrite_input(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Contact: Rashi Patil")
    document.save(source)
    with pytest.raises(ValueError, match="different"):
        redact_docx(source, source, detector=PIIDetector(use_spacy=False))
